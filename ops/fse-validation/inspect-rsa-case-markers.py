"""Read official RSA KO color markers without editing the workbooks."""
import hashlib
import json
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / 'ops/.local/fse-accreditamento/official-sources/d937255fd7e9c079c5641c537da17fe98a2f2259'


def fill_key(cell):
    fill = cell.fill
    if fill is None:
        return None
    color = fill.fgColor
    return (fill.patternType, color.type, str(color.value))


def main():
    path = SOURCE / 'RSA-KO.xlsx'
    manifest = json.loads((SOURCE / 'manifest.json').read_text(encoding='utf-8-sig'))
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    if digest != manifest['files']['RSA-KO.xlsx']['sha256']:
        raise RuntimeError('Official workbook hash changed')
    wb = load_workbook(path, read_only=True, data_only=False)
    legends = {}
    for row in wb['LEGENDA'].iter_rows():
        for cell in row:
            if cell.value:
                # In the official legend, text and its colored marker may be in adjacent cells.
                colors = {fill_key(c) for c in row if c.fill is not None and c.fill.patternType == 'solid'}
                for key in colors:
                    legends[key] = str(cell.value)
    print('Legend:', legends)
    result = {'source_sha256': digest, 'source_url': manifest['files']['RSA-KO.xlsx']['url'],
              'official_accreditation_evidence': False, 'cases': {}}
    for sheet in wb.worksheets:
        if sheet.title == 'LEGENDA':
            continue
        rows = []
        for row in sheet.iter_rows():
            marked = [c for c in row if fill_key(c) in legends]
            if marked:
                rows.append({'row': row[0].row, 'path': row[0].value,
                             'marked_cells': [{'cell': c.coordinate, 'value': c.value,
                                               'meaning': legends[fill_key(c)]} for c in marked],
                             'notes': row[4].value if len(row) > 4 else None})
        result['cases'][sheet.title] = rows
        print(sheet.title, json.dumps([{'row': r['row'], 'path': r['path'], 'notes': r['notes']} for r in rows if r['notes']], ensure_ascii=True))
    wb.close()
    output = SOURCE.parent.parent / 'rsa-ko-markers.json'
    output.write_text(json.dumps(result, indent=2, ensure_ascii=False, default=str), encoding='utf-8')


if __name__ == '__main__':
    main()
