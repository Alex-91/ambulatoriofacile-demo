"""Read-only extraction of official accreditation forms. Never writes a workbook."""
import json
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "ops/.local/fse-accreditamento/official-sources/d937255fd7e9c079c5641c537da17fe98a2f2259"
result = {}
for name in ("accreditamento-checklist_V9.0.0.xlsx", "RSA-OK.xlsx", "RSA-KO.xlsx"):
    wb = load_workbook(SOURCE / name, read_only=True, data_only=False)
    sheets = {}
    for ws in wb.worksheets:
        rows = []
        for index, row in enumerate(ws.iter_rows(), 1):
            cells = {cell.column_letter: cell.value for cell in row if cell.value is not None}
            if not cells:
                continue
            if name != "accreditamento-checklist_V9.0.0.xlsx" or index <= 10 or any("RSA" in str(v) for v in cells.values()):
                rows.append({"row": index, "cells": cells})
        sheets[ws.title] = rows
    result[name] = sheets
    wb.close()

ns = {"table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0", "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0"}
with ZipFile(SOURCE / "schedaAPIClient_HTTPS-FSE2_v3.3.ods") as z:
    xml = ET.fromstring(z.read("content.xml"))
    sheets = {}
    for table in xml.findall(".//table:table", ns):
        name = table.get("{" + ns["table"] + "}name")
        rows = []
        row_number = 1
        for row in table.findall("table:table-row", ns):
            values = []
            for cell in row:
                texts = ["".join(p.itertext()) for p in cell.findall("text:p", ns)]
                if texts:
                    values.append("\n".join(texts))
            if values:
                rows.append({"row": row_number, "values": values})
            row_number += int(row.get("{" + ns["table"] + "}number-rows-repeated", "1"))
        sheets[name] = rows
    result["schedaAPIClient_HTTPS-FSE2_v3.3.ods"] = sheets

destination = SOURCE.parent.parent / "official-forms-extracted.json"
destination.write_text(json.dumps(result, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
print(destination)
for name, sheets in result.items():
    print(name, {sheet: len(rows) for sheet, rows in sheets.items()})
